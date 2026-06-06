#!/usr/bin/env python3
"""Table 5: Scalability"""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

with open('analysis/outputs/aggregated/exp6_aggregated.json') as f:
    data = json.load(f)

wb = Workbook()
ws = wb.active
ws.title = "Scalability"

headers = ['Agents', 'Duration (ms)', 'Throughput (tps)']
ws.append(headers)

for col in range(1, len(headers) + 1):
    cell = ws.cell(1, col)
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    cell.alignment = Alignment(horizontal='center')

for d in data:
    ws.append([
        d['n_agents'],
        round(d['duration_mean'], 1),
        round(d['throughput_tps'], 1)
    ])

for col in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(col)].width = 15

wb.save('analysis/outputs/tables/table5_scalability.xlsx')
print('✓ table5_scalability.xlsx')
